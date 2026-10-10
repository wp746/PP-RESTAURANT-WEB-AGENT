from __future__ import annotations
import csv
import io
import json
import re
import shutil
import stat
import zipfile
from pathlib import Path, PurePosixPath

import yaml
from PIL import Image, ImageOps

TEXT_EXT = {'.txt', '.md', '.csv', '.json', '.yaml', '.yml'}
IMAGE_EXT = {'.jpg', '.jpeg', '.png', '.webp', '.avif'}
OFFICE_EXT = {'.docx', '.xlsx', '.pdf'}
ALLOWED_EXT = TEXT_EXT | IMAGE_EXT | OFFICE_EXT | {'.mp4', '.mov'}

class InputError(ValueError):
    pass


def extract_zip(archive: Path, destination: Path, max_entries=500, max_mb=350, max_file_mb=30):
    """Secure ZIP extraction: reject traversal/symlinks/bombs, never execute uploads."""
    destination.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(archive) as z:
        info = z.infolist()
        if len(info) > max_entries:
            raise InputError(f'Too many files ({len(info)} > {max_entries})')
        total = sum(item.file_size for item in info)
        if total > max_mb * 1024 * 1024:
            raise InputError('ZIP uncompressed size exceeds policy')
        for item in info:
            name = item.filename.replace('\\', '/')
            p = PurePosixPath(name)
            if (p.is_absolute() or '..' in p.parts or not name or '\x00' in name or
                (item.external_attr >> 16) & 0o170000 == stat.S_IFLNK):
                raise InputError(f'Unsafe path or symlink: {item.filename}')
            if item.file_size > max_file_mb * 1024 * 1024:
                raise InputError(f'File too large: {item.filename}')
            if item.is_dir():
                continue
            if Path(name).suffix.lower() not in ALLOWED_EXT:
                # Disallow active payloads and arbitrary uploaded code.
                continue
            target = destination.joinpath(*p.parts)
            target.parent.mkdir(parents=True, exist_ok=True)
            with z.open(item) as src, target.open('wb') as dst:
                shutil.copyfileobj(src, dst)
    return sorted(p for p in destination.rglob('*') if p.is_file())


def text_of(path: Path):
    ext = path.suffix.lower()
    if ext in {'.txt','.md','.csv','.json','.yaml','.yml'}:
        return path.read_text(encoding='utf-8-sig', errors='replace')[:60000]
    if ext == '.docx':
        try:
            from docx import Document
            d = Document(str(path))
            paras = [p.text for p in d.paragraphs]
            for t in d.tables:
                paras.extend(' | '.join(c.text for c in r.cells) for r in t.rows)
            return '\n'.join(paras)[:60000]
        except ImportError:
            return '[DOCX not parsed: install pp-restaurant-web-agent[files]]'
    if ext == '.xlsx':
        try:
            from openpyxl import load_workbook
            wb=load_workbook(path, read_only=True, data_only=True)
            rows=[]
            for sh in wb:
                rows.append(f'### {sh.title}')
                for r in sh.iter_rows(max_row=1200, values_only=True):
                    rows.append(' | '.join(str(x or '') for x in r[:12]))
            wb.close()
            return '\n'.join(rows)[:60000]
        except ImportError:
            return '[XLSX not parsed: install pp-restaurant-web-agent[files]]'
    if ext == '.pdf':
        try:
            from pypdf import PdfReader
            pdf=PdfReader(path)
            return '\n'.join((p.extract_text() or '') for p in pdf.pages[:35])[:60000]
        except ImportError:
            return '[PDF not parsed: install pp-restaurant-web-agent[files]]'
    return ''


def collect_text(files: list[Path], root: Path):
    docs = []
    for p in files:
        if p.suffix.lower() in TEXT_EXT | OFFICE_EXT:
            docs.append({'source':str(p.relative_to(root)), 'content':text_of(p)})
    return docs


def locate_merchant_manifest(files: list[Path]):
    """Optional authoritative manifest. Cannot infer approvals from ordinary prose."""
    exact = [p for p in files if p.name.lower() in {'merchant.json','merchant.yaml','merchant.yml'}]
    if not exact:
        return {}
    file=sorted(exact,key=lambda x:len(x.parts))[0]
    try:
        data=json.loads(file.read_text('utf-8-sig')) if file.suffix=='.json' else yaml.safe_load(file.read_text('utf-8-sig'))
        if not isinstance(data,dict):raise ValueError('manifest must be an object')
        return data
    except (ValueError, yaml.YAMLError) as e:
        raise InputError(f'Invalid merchant manifest: {e}')


def optimize_images(files: list[Path], output: Path, max_width=1920, quality=80):
    output.mkdir(parents=True, exist_ok=True)
    photos=[]
    for i,src in enumerate(f for f in files if f.suffix.lower() in IMAGE_EXT):
        try:
            with Image.open(src) as img:
                img=ImageOps.exif_transpose(img)
                img.thumbnail((max_width,max_width*2),Image.Resampling.LANCZOS)
                transparent='A' in img.getbands() and img.getextrema()[-1][0] < 255
                img=img.convert('RGBA' if transparent else 'RGB')
                name=f'image-{i+1:03}.webp'
                img.save(output/name,'WEBP',quality=quality,method=6)
                src_key=str(src).lower()
                if any(w in src_key for w in ('菜品','产品','主菜','招牌','food','dish','menu-photo')):kind='product'
                elif any(w in src_key for w in ('卡座','包厢','大厅','座位','环境','interior','dining')):kind='interior'
                elif any(w in src_key for w in ('门头','外景','招牌门','storefront','exterior')):kind='storefront'
                elif any(w in src_key for w in ('工艺','后厨','烹饪','制作','process')):kind='process'
                else:kind='unknown'
                photos.append({'path':f'assets/{name}','source':str(src),'kind_hint':kind,'original_filename':src.name,'width':img.width,'height':img.height,'bytes':(output/name).stat().st_size})
        except Exception:
            continue
    return photos


def extract_basic_facts(docs:list[dict], manifest:dict):
    """Conservative: no invented field values; manifest has priority."""
    mapping = {
      'name':['店名','门店名称','品牌名称'],
      'category':['主营品类','餐饮品类','菜系','品类'],
      'address':['门店地址','详细地址','地址'],
      'phone':['联系电话','联系电话号码','商家电话','电话'],
      'opening_hours':['营业时间','营业时段'],
      'story':['品牌故事','创办故事','开店故事'],
      'slogan':['品牌口号','品牌主张','slogan'],
    }
    facts={}
    evidence={}
    for field,labels in mapping.items():
        value=manifest.get(field)
        if isinstance(value,str) and value.strip():
            facts[field]=value.strip();evidence[field]='merchant.yaml / merchant.json';continue
        for doc in docs:
            for line in doc['content'].splitlines():
                for label in labels:
                    m=re.match(r'^\s*(?:[-*]\s*)?'+re.escape(label)+r'\s*[：:]\s*(.+?)\s*$',line,re.I)
                    if m and m.group(1).strip() not in {'无','暂无','待补','—','--','-','/'}:
                        facts[field]=m.group(1).strip()[:500]
                        evidence[field]=doc['source']
                        break
                if field in facts:break
            if field in facts:break
    # canonical record may include additional user supplied facts, not inferred history
    for k in ('menu','offers','map_url','official_url','contact_url','culture','service','audience','price_band'):
        if manifest.get(k):facts[k]=manifest[k];evidence[k]='merchant manifest'
    return facts,evidence
